from __future__ import annotations

import tempfile
import unittest
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from scripts.weekly_dashboard.core import (
    build_history_row,
    load_config,
    parse_management_pdf_text,
    read_ads,
    read_lstep,
    read_management,
    target_week,
    update_history,
    write_outputs,
)


CONFIG = {
    "column_mapping": {"added_at": "友だち追加日時"},
    "weekly_dashboard": {
        "lstep": {"header_row": 1, "encodings": ["utf-8-sig"]},
        "ads": {
            "encodings": ["utf-8-sig"],
            "separators": [","],
            "header_search_rows": 3,
            "column_aliases": {
                "date": ["日"],
                "clicks": ["クリック数"],
            },
        },
    },
    "status_tags": {"paid": "STD_決済済み"},
    "purchase_tags": {},
    "tags": {},
    "branch_tags": {},
}


class WeekTests(unittest.TestCase):
    def test_tuesday_run_targets_previous_monday_to_sunday(self) -> None:
        week = target_week(date(2026, 7, 14))
        self.assertEqual(date(2026, 7, 6), week.start)
        self.assertEqual(date(2026, 7, 12), week.end)

    def test_week_cut_is_stable_for_other_run_weekdays(self) -> None:
        week = target_week(date(2026, 7, 15))
        self.assertEqual(date(2026, 7, 6), week.start)
        self.assertEqual(date(2026, 7, 12), week.end)


class LstepTests(unittest.TestCase):
    def test_full_and_prefiltered_exports_have_same_weekly_count(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            full = root / "full.csv"
            filtered = root / "filtered.csv"
            pd.DataFrame(
                [
                    {"友だち追加日時": "2026-07-05 23:59", "STD_決済済み": 1, "表示名": "discard"},
                    {"友だち追加日時": "2026-07-06 00:00", "STD_決済済み": 0, "表示名": "discard"},
                    {"友だち追加日時": "2026-07-12 23:59", "STD_決済済み": 1, "表示名": "discard"},
                    {"友だち追加日時": "2026-07-13 00:00", "STD_決済済み": 1, "表示名": "discard"},
                ]
            ).to_csv(full, index=False, encoding="utf-8-sig")
            pd.DataFrame(
                [
                    {"友だち追加日時": "2026-07-06 00:00", "STD_決済済み": 0},
                    {"友だち追加日時": "2026-07-12 23:59", "STD_決済済み": 1},
                ]
            ).to_csv(filtered, index=False, encoding="utf-8-sig")
            week = target_week(date(2026, 7, 14))

            full_count, _ = read_lstep(full, CONFIG, week)
            filtered_count, _ = read_lstep(filtered, CONFIG, week)
            self.assertEqual(2, full_count)
            self.assertEqual(full_count, filtered_count)


class AdsTests(unittest.TestCase):
    def test_ads_clicks_are_filtered_to_target_week(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "ads.csv"
            path.write_text(
                "Google広告レポート\n"
                "日,クリック数\n"
                "2026-07-05,99\n"
                "2026-07-06,10\n"
                "2026-07-12,20\n"
                "2026-07-13,88\n",
                encoding="utf-8-sig",
            )
            clicks = read_ads(path, CONFIG, target_week(date(2026, 7, 14)))
            self.assertEqual(30, clicks)


class ManagementPdfTests(unittest.TestCase):
    def test_pdf_row_parser_handles_concatenated_impressions_clicks(self) -> None:
        text = (
            "2026/07/06月"
            "1,000 2,000 200% 10,000 500 29843823 2.76% 0 10 1% 50 "
            "5 50% 1 2 3 4 5"
        )
        frame = parse_management_pdf_text(text)
        self.assertEqual(1, len(frame))
        self.assertEqual(823, int(frame.iloc[0]["clicks"]))
        self.assertEqual(10, int(frame.iloc[0]["cv_f"]))
        self.assertEqual(5, int(frame.iloc[0]["monshin_answers"]))
        self.assertEqual(15, int(frame.iloc[0]["purchase_cv"]))


class HistoryTests(unittest.TestCase):
    def test_history_upsert_is_idempotent_and_preserves_prior_week(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "funnel_history.csv"
            first_week = target_week(date(2026, 7, 7))
            second_week = target_week(date(2026, 7, 14))
            first = build_history_row(
                "std",
                first_week,
                {"clicks": 100, "cv_f": 10, "monshin_answers": 8, "purchase_cv": 4},
                20,
            )
            second = build_history_row(
                "std",
                second_week,
                {"clicks": 120, "cv_f": 12, "monshin_answers": 9, "purchase_cv": 5},
                24,
            )
            update_history(path, first)
            update_history(path, second)
            history = update_history(path, second)
            self.assertEqual(2, len(history))
            self.assertEqual(
                [first_week.start.isoformat(), second_week.start.isoformat()],
                history["week_start"].tolist(),
            )


class IntegrationTests(unittest.TestCase):
    def test_rerun_is_idempotent_and_never_generates_markdown_report(self) -> None:
        repo_root = Path(__file__).resolve().parents[3]
        config = load_config(repo_root, "std")
        week = target_week(date(2026, 7, 14))
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            inputs = root / "raw"
            outputs = root / "output"
            inputs.mkdir()
            (inputs / "ads.csv").write_text(
                "report\nperiod\n日\t検索語句\tクリック数\n"
                "2026-07-06\t合計: アカウント\t10\n"
                "2026-07-12\t合計: アカウント\t20\n",
                encoding="utf-16",
            )
            (inputs / "lstep.csv").write_text(
                "internal_added,internal_paid\n"
                "友だち追加日時,STD_決済済み\n"
                "2026-07-06 10:00,0\n"
                "2026-07-12 10:00,1\n",
                encoding="cp932",
            )
            pd.DataFrame(
                [
                    {"日付": "2026-07-06", "CV(F)": 3, "問診回答": 2, "購入CV": 1},
                    {"日付": "2026-07-12", "CV(F)": 4, "問診回答": 3, "購入CV": 2},
                ]
            ).to_csv(inputs / "management.csv", index=False, encoding="utf-8-sig")

            line_count, _ = read_lstep(inputs / "lstep.csv", config, week)
            clicks = read_ads(inputs / "ads.csv", config, week)
            management = read_management(inputs, config, week)
            management["clicks"] = clicks
            row = build_history_row("std", week, management, line_count)
            history_path = root / "funnel_history.csv"

            first = update_history(history_path, row)
            write_outputs(outputs, config, row, first, sorted(inputs.iterdir()))
            second = update_history(history_path, row)
            write_outputs(outputs, config, row, second, sorted(inputs.iterdir()))

            self.assertEqual(1, len(second))
            self.assertFalse((outputs / "weekly_report.md").exists())
            self.assertEqual(
                ["Weekly"],
                load_workbook(outputs / "weekly_dashboard.xlsx", read_only=True).sheetnames,
            )
            self.assertEqual(
                {"weekly_dashboard.xlsx", "weekly_kpi.csv", "bottlenecks.csv",
                 "analysis_meta.json", "funnel_history.csv"},
                {path.name for path in outputs.iterdir()},
            )


if __name__ == "__main__":
    unittest.main()
